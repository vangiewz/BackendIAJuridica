"""Excel con openpyxl: la hoja de datos y una hoja con la ficha del reporte."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.reportes.esquemas import ReporteResultado
from app.services.reportes.exportadores import comun

_RELLENO = PatternFill("solid", fgColor="1B6B4A")
_ANCHO_MAXIMO = 60


def _hoja_reporte(libro: Workbook, reporte: ReporteResultado) -> None:
    hoja = libro.active
    hoja.title = "Reporte"
    cabeceras = comun.encabezados(reporte)
    filas = comun.filas_de_texto(reporte)

    hoja.append([reporte.titulo])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append([f"Generado el {comun.generado_en()} · {reporte.total} resultados"])
    hoja.append([])

    fila_cabecera = hoja.max_row + 1
    hoja.append(cabeceras)
    for celda in hoja[fila_cabecera]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _RELLENO
        celda.alignment = Alignment(vertical="center", wrap_text=True)

    for fila in filas:
        hoja.append(fila)

    for indice, cabecera in enumerate(cabeceras, start=1):
        largos = [len(cabecera)] + [len(fila[indice - 1]) for fila in filas]
        hoja.column_dimensions[get_column_letter(indice)].width = min(
            max(12, max(largos) + 2), _ANCHO_MAXIMO)
    # Los encabezados quedan a la vista al desplazarse por los datos.
    hoja.freeze_panes = hoja.cell(row=fila_cabecera + 1, column=1)

    _agregar_grafico(hoja, reporte, fila_cabecera, len(filas))


def _agregar_grafico(hoja, reporte: ReporteResultado, fila_cabecera: int, total: int) -> None:
    """Gráfico nativo de Excel sobre las mismas celdas, no una imagen pegada."""
    if reporte.visualizacion not in ("barras", "torta") or not total:
        return
    ejes = comun.datos_de_grafico(reporte)
    if ejes is None:
        return
    columna_etiqueta = next(i for i, c in enumerate(reporte.columnas, start=1)
                            if c.tipo != "numero")
    columna_valor = next(i for i, c in enumerate(reporte.columnas, start=1)
                         if c.tipo == "numero")
    etiquetas = Reference(hoja, min_col=columna_etiqueta, min_row=fila_cabecera + 1,
                          max_row=fila_cabecera + total)
    valores = Reference(hoja, min_col=columna_valor, min_row=fila_cabecera,
                        max_row=fila_cabecera + total)

    grafico = BarChart() if reporte.visualizacion == "barras" else PieChart()
    grafico.title = reporte.titulo
    grafico.add_data(valores, titles_from_data=True)
    grafico.set_categories(etiquetas)
    grafico.height, grafico.width = 9, 18
    hoja.add_chart(grafico, f"A{fila_cabecera + total + 3}")


def _hoja_metadatos(libro: Workbook, reporte: ReporteResultado) -> None:
    hoja = libro.create_sheet("Metadatos")
    hoja.append(["Dato", "Valor"])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _RELLENO
    for clave, valor in comun.metadatos(reporte):
        hoja.append([clave, valor])
    hoja.column_dimensions["A"].width = 22
    hoja.column_dimensions["B"].width = _ANCHO_MAXIMO
    for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=2):
        fila[0].alignment = Alignment(wrap_text=True, vertical="top")


def generar(reporte: ReporteResultado) -> bytes:
    libro = Workbook()
    _hoja_reporte(libro, reporte)
    _hoja_metadatos(libro, reporte)
    memoria = BytesIO()
    libro.save(memoria)
    return memoria.getvalue()
