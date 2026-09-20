"""Word con python-docx: título, ficha del reporte y tabla real de Word."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models.reportes.esquemas import ReporteResultado
from app.services.reportes.exportadores import comun

ESTILO_TABLA = "Light Grid Accent 1"
ANCHO_BARRA = 28


def _ficha(documento: Document, reporte: ReporteResultado) -> None:
    for clave, valor in comun.metadatos(reporte):
        parrafo = documento.add_paragraph()
        corrida = parrafo.add_run(f"{clave}: ")
        corrida.bold = True
        parrafo.add_run(valor)
        parrafo.paragraph_format.space_after = Pt(2)


def _tabla(documento: Document, reporte: ReporteResultado) -> None:
    cabeceras = comun.encabezados(reporte)
    filas = comun.filas_de_texto(reporte)
    tabla = documento.add_table(rows=1, cols=len(cabeceras))
    try:
        tabla.style = ESTILO_TABLA
    except KeyError:
        # La plantilla por defecto puede no traer ese estilo; la tabla sigue siendo válida.
        tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, cabeceras):
        celda.text = ""
        corrida = celda.paragraphs[0].add_run(texto)
        corrida.bold = True
    for fila in filas:
        celdas = tabla.add_row().cells
        for celda, texto in zip(celdas, fila):
            celda.text = texto


def _grafico_de_texto(documento: Document, reporte: ReporteResultado) -> None:
    """El gráfico como barras de bloques: es proporcional y no necesita una imagen."""
    ejes = comun.datos_de_grafico(reporte)
    if ejes is None:
        _tabla(documento, reporte)
        return
    etiquetas, valores, titulo_valor = ejes
    maximo = max(valores) if valores else 0
    total = sum(valores)
    documento.add_paragraph(titulo_valor).runs[0].bold = True

    tabla = documento.add_table(rows=1, cols=3)
    tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, ("Categoría", "Proporción", titulo_valor)):
        corrida = celda.paragraphs[0].add_run(texto)
        corrida.bold = True
    for etiqueta, valor in zip(etiquetas, valores):
        celdas = tabla.add_row().cells
        celdas[0].text = etiqueta
        bloques = int(round((valor / maximo) * ANCHO_BARRA)) if maximo else 0
        celdas[1].text = "█" * max(bloques, 1 if valor else 0)
        porcentaje = f" ({valor / total * 100:.1f}%)" if total else ""
        celdas[2].text = f"{valor:g}{porcentaje}"


def _resumen(documento: Document, reporte: ReporteResultado) -> None:
    for fila in reporte.filas:
        for columna in reporte.columnas:
            parrafo = documento.add_paragraph()
            cifra = parrafo.add_run(comun.formatear(fila.get(columna.clave), columna.tipo))
            cifra.bold = True
            cifra.font.size = Pt(20)
            parrafo.add_run(f"   {columna.etiqueta}")


def generar(reporte: ReporteResultado) -> bytes:
    documento = Document()
    documento.add_heading(reporte.titulo, level=1)
    subtitulo = documento.add_paragraph(f"Generado el {comun.generado_en()}")
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.LEFT

    documento.add_heading("Datos del reporte", level=2)
    _ficha(documento, reporte)

    documento.add_heading("Resultados", level=2)
    if reporte.total == 0:
        documento.add_paragraph("No hay datos que cumplan esas condiciones.")
    elif reporte.visualizacion == "resumen":
        _resumen(documento, reporte)
    elif reporte.visualizacion in ("barras", "torta"):
        _grafico_de_texto(documento, reporte)
    else:
        _tabla(documento, reporte)

    memoria = BytesIO()
    documento.save(memoria)
    return memoria.getvalue()
