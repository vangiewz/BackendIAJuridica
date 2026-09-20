"""PowerPoint con python-pptx: las diapositivas salen del reporte, no de una plantilla fija.

Cuántas diapositivas hay y qué contienen depende de lo que devolvió la consulta: un
gráfico si el reporte es un gráfico, indicadores si es un resumen, y la tabla partida en
varias diapositivas si es un listado. Se corta en un máximo razonable para no producir
un archivo de cien diapositivas.
"""
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches, Pt

from app.models.reportes.esquemas import ReporteResultado
from app.services.reportes.exportadores import comun

FILAS_POR_DIAPOSITIVA = 10
MAXIMO_DIAPOSITIVAS_DE_TABLA = 10
VERDE = RGBColor(0x1B, 0x6B, 0x4A)
TINTA = RGBColor(0x22, 0x20, 0x1C)
SUAVE = RGBColor(0x5F, 0x59, 0x52)

DISENO_TITULO, DISENO_TITULO_Y_CUERPO, DISENO_SOLO_TITULO, DISENO_VACIO = 0, 1, 5, 6


def _texto(marco, lineas, tamano=14, color=TINTA):
    marco.word_wrap = True
    for indice, linea in enumerate(lineas):
        parrafo = marco.paragraphs[0] if indice == 0 else marco.add_paragraph()
        parrafo.text = linea
        for corrida in parrafo.runs:
            corrida.font.size = Pt(tamano)
            corrida.font.color.rgb = color


def _portada(presentacion, reporte: ReporteResultado) -> None:
    diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[DISENO_TITULO])
    diapositiva.shapes.title.text = reporte.titulo
    subtitulo = diapositiva.placeholders[1]
    lineas = [reporte.peticion] if reporte.peticion else []
    lineas.append(f"{reporte.entidad_etiqueta} · {reporte.total} resultados")
    lineas.append(f"Generado el {comun.generado_en()}")
    _texto(subtitulo.text_frame, lineas, tamano=14, color=SUAVE)


def _ficha(presentacion, reporte: ReporteResultado) -> None:
    diapositiva = presentacion.slides.add_slide(
        presentacion.slide_layouts[DISENO_TITULO_Y_CUERPO])
    diapositiva.shapes.title.text = "Datos del reporte"
    lineas = [f"{clave}: {valor}" for clave, valor in comun.metadatos(reporte)
              if clave != "Petición"]
    _texto(diapositiva.placeholders[1].text_frame, lineas, tamano=13)


def _diapositiva_de_grafico(presentacion, reporte: ReporteResultado) -> bool:
    """Gráfico nativo de PowerPoint, editable, construido con los datos reales."""
    ejes = comun.datos_de_grafico(reporte)
    if ejes is None:
        return False
    etiquetas, valores, titulo_valor = ejes
    diapositiva = presentacion.slides.add_slide(
        presentacion.slide_layouts[DISENO_SOLO_TITULO])
    diapositiva.shapes.title.text = reporte.titulo

    datos = CategoryChartData()
    datos.categories = etiquetas
    datos.add_series(titulo_valor, valores)
    tipo = (XL_CHART_TYPE.PIE if reporte.visualizacion == "torta"
            else XL_CHART_TYPE.COLUMN_CLUSTERED)
    marco = diapositiva.shapes.add_chart(
        tipo, Inches(0.8), Inches(1.6), Inches(8.4), Inches(4.8), datos)
    grafico = marco.chart
    grafico.has_legend = True
    grafico.legend.position = XL_LEGEND_POSITION.BOTTOM
    grafico.legend.include_in_layout = False
    return True


def _diapositivas_de_tabla(presentacion, reporte: ReporteResultado) -> None:
    cabeceras = comun.encabezados(reporte)
    filas = comun.filas_de_texto(reporte)
    bloques = [filas[i:i + FILAS_POR_DIAPOSITIVA]
               for i in range(0, len(filas), FILAS_POR_DIAPOSITIVA)]
    recortado = len(bloques) > MAXIMO_DIAPOSITIVAS_DE_TABLA
    bloques = bloques[:MAXIMO_DIAPOSITIVAS_DE_TABLA]

    for indice, bloque in enumerate(bloques, start=1):
        diapositiva = presentacion.slides.add_slide(
            presentacion.slide_layouts[DISENO_SOLO_TITULO])
        sufijo = f" ({indice}/{len(bloques)})" if len(bloques) > 1 else ""
        diapositiva.shapes.title.text = f"Resultados{sufijo}"
        forma = diapositiva.shapes.add_table(
            len(bloque) + 1, len(cabeceras),
            Inches(0.4), Inches(1.5), Inches(9.2), Inches(0.4 * (len(bloque) + 1)))
        tabla = forma.table
        for columna, texto in enumerate(cabeceras):
            celda = tabla.cell(0, columna)
            celda.text = texto
            for parrafo in celda.text_frame.paragraphs:
                for corrida in parrafo.runs:
                    corrida.font.size = Pt(11)
                    corrida.font.bold = True
        for fila_indice, fila in enumerate(bloque, start=1):
            for columna, texto in enumerate(fila):
                celda = tabla.cell(fila_indice, columna)
                celda.text = texto
                for parrafo in celda.text_frame.paragraphs:
                    for corrida in parrafo.runs:
                        corrida.font.size = Pt(10)

    if recortado:
        diapositiva = presentacion.slides.add_slide(
            presentacion.slide_layouts[DISENO_TITULO_Y_CUERPO])
        diapositiva.shapes.title.text = "Resultados restantes"
        mostradas = MAXIMO_DIAPOSITIVAS_DE_TABLA * FILAS_POR_DIAPOSITIVA
        _texto(diapositiva.placeholders[1].text_frame, [
            f"Se incluyeron las primeras {mostradas} filas de {len(filas)}.",
            "El archivo Excel del mismo reporte trae todas las filas.",
        ], tamano=14, color=SUAVE)


def _diapositiva_de_indicadores(presentacion, reporte: ReporteResultado) -> None:
    diapositiva = presentacion.slides.add_slide(
        presentacion.slide_layouts[DISENO_SOLO_TITULO])
    diapositiva.shapes.title.text = reporte.titulo
    pares = [(columna.etiqueta, comun.formatear(fila.get(columna.clave), columna.tipo))
             for fila in reporte.filas for columna in reporte.columnas]
    ancho = Inches(9.2 / max(min(len(pares), 4), 1))
    for indice, (etiqueta, valor) in enumerate(pares[:4]):
        izquierda = Emu(int(Inches(0.4)) + int(ancho) * indice)
        caja = diapositiva.shapes.add_textbox(izquierda, Inches(2.2), ancho, Inches(2))
        marco = caja.text_frame
        marco.word_wrap = True
        marco.paragraphs[0].text = valor
        marco.paragraphs[0].runs[0].font.size = Pt(40)
        marco.paragraphs[0].runs[0].font.bold = True
        marco.paragraphs[0].runs[0].font.color.rgb = VERDE
        parrafo = marco.add_paragraph()
        parrafo.text = etiqueta
        parrafo.runs[0].font.size = Pt(13)
        parrafo.runs[0].font.color.rgb = SUAVE


def generar(reporte: ReporteResultado) -> bytes:
    presentacion = Presentation()
    _portada(presentacion, reporte)
    _ficha(presentacion, reporte)

    if reporte.total == 0:
        diapositiva = presentacion.slides.add_slide(
            presentacion.slide_layouts[DISENO_TITULO_Y_CUERPO])
        diapositiva.shapes.title.text = "Sin resultados"
        _texto(diapositiva.placeholders[1].text_frame,
               ["No hay datos que cumplan esas condiciones."], color=SUAVE)
    elif reporte.visualizacion == "resumen":
        _diapositiva_de_indicadores(presentacion, reporte)
    elif reporte.visualizacion in ("barras", "torta"):
        # El gráfico primero y la tabla detrás: si el gráfico no se puede armar, la
        # tabla sigue diciendo lo mismo y no se pierde ningún dato.
        _diapositiva_de_grafico(presentacion, reporte)
        _diapositivas_de_tabla(presentacion, reporte)
    else:
        _diapositivas_de_tabla(presentacion, reporte)

    memoria = BytesIO()
    presentacion.save(memoria)
    return memoria.getvalue()
