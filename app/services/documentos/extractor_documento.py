import io
import re
from pathlib import Path
import pypdf
import docx

class FormatoNoSoportadoError(Exception): ...
class DocumentoSinTextoError(Exception):
    """El archivo se abrio pero no tiene capa de texto (tipicamente, un PDF escaneado)."""

EXTENSIONES = {".pdf", ".docx", ".txt"}
MINIMO_CARACTERES = 200

def extraer_texto(contenido: bytes, nombre_archivo: str) -> str:
    """Texto plano del documento. El formato sale de la extension del nombre."""
    ext = Path(nombre_archivo).suffix.lower()
    
    if ext not in EXTENSIONES:
        raise FormatoNoSoportadoError(f"Extension {ext} no soportada.")
        
    archivo_io = io.BytesIO(contenido)
    texto_crudo = ""
    
    if ext == ".pdf":
        reader = pypdf.PdfReader(archivo_io)
        paginas = []
        for page in reader.pages:
            texto_pagina = page.extract_text()
            if texto_pagina:
                paginas.append(texto_pagina)
        texto_crudo = "\n".join(paginas)
        
    elif ext == ".docx":
        doc = docx.Document(archivo_io)
        partes = []
        # Se recorre el XML del cuerpo en vez de `doc.paragraphs` + `doc.tables` para que los
        # parrafos y las tablas salgan en el orden en que estan escritos. En una minuta los
        # datos de las partes suelen ir en una tabla al principio, y concatenar todas las
        # tablas al final desordenaria el documento.
        for child in doc.element.body:
            if child.tag.endswith('p'):
                para = docx.text.paragraph.Paragraph(child, doc)
                if para.text:
                    partes.append(para.text)
            elif child.tag.endswith('tbl'):
                table = docx.table.Table(child, doc)
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            partes.append(cell.text)
        texto_crudo = "\n".join(partes)
        
    elif ext == ".txt":
        try:
            texto_crudo = contenido.decode("utf-8")
        except UnicodeDecodeError:
            texto_crudo = contenido.decode("latin-1")
            
    # Normalizar saltos de linea
    texto_norm = texto_crudo.replace("\r\n", "\n")
    texto_norm = re.sub(r'\n{3,}', '\n\n', texto_norm)
    
    # Contar caracteres no vacíos
    caracteres_validos = len("".join(texto_norm.split()))
    
    if caracteres_validos < MINIMO_CARACTERES:
        raise DocumentoSinTextoError("Documento sin suficiente texto válido.")
        
    return texto_norm
