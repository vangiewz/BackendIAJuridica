import io
import pytest
from pathlib import Path
import docx
import pypdf
from app.services.documentos.extractor_documento import (
    extraer_texto, 
    FormatoNoSoportadoError, 
    DocumentoSinTextoError
)

def test_extraer_texto_pdf_real():
    ruta_pdf = Path(__file__).parent.parent / "data" / "normativa" / "codigo_civil_oea.pdf"
    with open(ruta_pdf, "rb") as f:
        contenido = f.read()
    
    texto = extraer_texto(contenido, "codigo.pdf")
    assert len(texto) > 200
    assert "ARTÍCULO" in texto or "ARTICULO" in texto

def test_pdf_pagina_blanca_lanza_error():
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    
    archivo_io = io.BytesIO()
    writer.write(archivo_io)
    contenido = archivo_io.getvalue()
    
    with pytest.raises(DocumentoSinTextoError):
        extraer_texto(contenido, "blanco.pdf")

def test_docx_parrafos_y_tabla():
    doc = docx.Document()
    doc.add_paragraph("Primer párrafo de prueba. " + "Este es un texto de relleno para superar la restricción de doscientos caracteres mínimos. " * 3)
    
    tabla = doc.add_table(rows=1, cols=2)
    celdas = tabla.rows[0].cells
    celdas[0].text = "Celda1"
    celdas[1].text = "Celda2"
    
    doc.add_paragraph("Segundo párrafo final.")
    
    archivo_io = io.BytesIO()
    doc.save(archivo_io)
    contenido = archivo_io.getvalue()
    
    texto = extraer_texto(contenido, "documento.docx")
    
    assert "Primer párrafo de prueba." in texto
    assert "Celda1" in texto
    assert "Celda2" in texto
    assert "Segundo párrafo final." in texto

def test_formato_no_soportado():
    with pytest.raises(FormatoNoSoportadoError):
        extraer_texto(b"fake_content", "archivo.xlsx")

def test_txt_latin1_con_acentos():
    texto_original = "Este documento tiene acentos: canción, árbol."
    # Primero vemos si con latin-1 pasa.
    # Lo rellenamos para pasar el minimo de caracteres.
    texto_largo = (texto_original + " ") * 20
    contenido_latin1 = texto_largo.encode("latin-1")
    
    # Esto no debe lanzar excepcion y debe decodificar bien
    texto = extraer_texto(contenido_latin1, "documento.txt")
    assert "canción" in texto
    assert "árbol" in texto

def test_normalizacion_de_saltos():
    # Rellenamos con texto basura para superar el limite de 200 caracteres
    texto_basura = "Este es un texto de relleno para superar el límite de doscientos caracteres impuesto por la extracción de texto. " * 5
    
    texto_crudo = "Linea 1\r\nLinea 2\n\n\n\nLinea 3\r\n\r\n\r\nLinea 4" + "\n" + texto_basura
    contenido = texto_crudo.encode("utf-8")
    
    texto_extraido = extraer_texto(contenido, "saltos.txt")
    
    assert "\r\n" not in texto_extraido
    assert "\n\n\n" not in texto_extraido
    assert "Linea 1\nLinea 2\n\nLinea 3\n\nLinea 4" in texto_extraido

